#!/usr/bin/env python3
"""Build editable CBC manuscript, cover letter and supplementary-information DOCX files."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
# Standalone public release keeps the manuscript files under manuscript/; the workspace under results/.
OUTDIR = ROOT / "manuscript" if (ROOT / "manuscript" / "source_data").is_dir() else ROOT / "results/manuscript_computational_biology_chemistry"
MAIN_MD = OUTDIR / "manuscript_CBC_en.md"
COVER_MD = OUTDIR / "cover_letter_CBC_en.md"
SI_MD = OUTDIR / "supplementary_information_CBC_en.md"
FIGDIR = OUTDIR / "figures"


def set_borders(element, **edges):
    # CBC tables: horizontal rules only (no vertical rules, no shading).
    borders = OxmlElement("w:tblBorders" if element.tag == qn("w:tblPr") else "w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in edges and element.tag != qn("w:tblPr"):
            continue
        el = OxmlElement(f"w:{edge}")
        style = edges.get(edge)
        el.set(qn("w:val"), style or "nil")
        if style:
            el.set(qn("w:sz"), "8"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")
        borders.append(el)
    element.append(borders)


def default_style(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_after = Pt(5)
    for name, size, color in (("Title", 16, "000000"), ("Heading 1", 13, "000000"), ("Heading 2", 11.5, "000000"), ("Heading 3", 10.5, "000000")):
        st = styles[name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(9)
        st.paragraph_format.space_after = Pt(5)
    if "Caption Compact" not in [s.name for s in styles]:
        st = styles.add_style("Caption Compact", WD_STYLE_TYPE.PARAGRAPH)
    else:
        st = styles["Caption Compact"]
    st.font.name = "Arial"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.font.size = Pt(8.5)
    st.paragraph_format.space_after = Pt(5)


def layout(doc):
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(3)
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    for typ, val in (("begin", "begin"), ("instrText", " PAGE "), ("end", "end")):
        el = OxmlElement("w:fldChar" if typ != "instrText" else "w:instrText")
        if typ == "instrText":
            el.set(qn("xml:space"), "preserve")
            el.text = val
        else:
            el.set(qn("w:fldCharType"), val)
        r._r.append(el)


def inline(p, text):
    token = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
    pos = 0
    for m in token.finditer(text):
        if m.start() > pos:
            p.add_run(text[pos:m.start()])
        t = m.group(0)
        if t.startswith("**"):
            r = p.add_run(t[2:-2]); r.bold = True
        elif t.startswith("*"):
            r = p.add_run(t[1:-1]); r.italic = True
        else:
            r = p.add_run(t[1:-1]); r.font.name = "Courier New"
        pos = m.end()
    if pos < len(text):
        p.add_run(text[pos:])


def add_para(doc, text, style="Normal"):
    p = doc.add_paragraph(style=style)
    inline(p, text)
    return p


def add_table(doc, rows):
    if not rows:
        return
    data = []
    for row in rows:
        if set(row.replace("|", "").strip()) <= {"-", ":", " "}:
            continue
        data.append([x.strip() for x in row.strip().strip("|").split("|")])
    if not data:
        return
    ncol = max(len(x) for x in data)
    table = doc.add_table(rows=len(data), cols=ncol)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_borders(table._tbl.tblPr, top="single", bottom="single")
    for i, row in enumerate(data):
        tr_pr = table.rows[i]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if i == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
        for j in range(ncol):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if j < len(row):
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.0
                inline(p, row[j])
            if i == 0:
                set_borders(cell._tc.get_or_add_tcPr(), bottom="single")
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    doc.add_paragraph()


def add_title_page(doc):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Paired-cohort meta-analysis with dual RNA-seq replication of recurrent transcriptional alterations in hepatocellular carcinoma")
    r.bold = True; r.font.size = Pt(16); r.font.color.rgb = RGBColor(0, 0, 0)
    p = doc.add_paragraph("Research Article"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Affiliations use lower-case superscript letters (CBC title-page rule).
    p = doc.add_paragraph()
    for text, sup in (("Simiao Yu", False), ("a", True), ("; Zhouping Wang", False), ("a,*", True)):
        p.add_run(text).font.superscript = sup
    p = doc.add_paragraph()
    p.add_run("a").font.superscript = True
    p.add_run(" State Key Laboratory of Food Science and Resources, School of Food Science and Technology, Jiangnan University, No. 1800 Lihu Avenue, Binhu District, Wuxi, Jiangsu 214122, China")
    for text in [
        "*Corresponding author: Zhouping Wang; wangzp@jiangnan.edu.cn; +86-510-85326195",
        "Running title: Recurrent HCC transcriptional alterations",
        "Clinical trial number: not applicable.",
    ]:
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if text == "Running title: Recurrent HCC transcriptional alterations" else WD_ALIGN_PARAGRAPH.LEFT
    doc.add_page_break()


def add_image(doc, path, width=6.4):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))


def parse_md(doc, path, include_figures=True):
    lines = path.read_text(encoding="utf-8").splitlines()
    i = 0
    skip_preamble = path.name == MAIN_MD.name
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1; continue
        if skip_preamble:
            if s == "## Abstract":
                skip_preamble = False
            else:
                i += 1; continue
        if s.startswith("# "):
            # The document builder supplies the title page/section title.
            i += 1; continue
        if s.startswith("## "):
            doc.add_heading(s[3:], level=1); i += 1; continue
        if s.startswith("### "):
            doc.add_heading(s[4:], level=2); i += 1; continue
        if s.startswith("#### "):
            doc.add_heading(s[5:], level=3); i += 1; continue
        if s.startswith("![") and include_figures:
            m = re.search(r"\(([^)]+)\)", s)
            if m:
                rel = m.group(1)
                pth = (path.parent / rel).resolve()
                if pth.exists(): add_image(doc, pth, 6.25)
            i += 1; continue
        if s.startswith(("**Figure", "**Supplementary Figure", "**Table", "**Graphical abstract", "*Notes:*")):
            add_para(doc, s, style="Caption Compact"); i += 1; continue
        if s.startswith("|"):
            rows=[]
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip()); i += 1
            add_table(doc, rows); continue
        if re.match(r"^\d+\.\s", s):
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.2); p.paragraph_format.first_line_indent=Inches(-.2); inline(p,s); i += 1; continue
        if s.startswith("**") and s.endswith("**"):
            add_para(doc,s); i += 1; continue
        add_para(doc,s); i += 1


def build_main():
    # Figures are uploaded as separate files; their captions sit before the references.
    doc=Document(); default_style(doc); layout(doc); add_title_page(doc); parse_md(doc, MAIN_MD, include_figures=False)
    out=OUTDIR/"manuscript_CBC_submission.docx"; doc.save(out); return out


def build_cover():
    doc=Document(); default_style(doc); layout(doc)
    doc.add_heading("Cover letter", level=1); parse_md(doc,COVER_MD,include_figures=False)
    out=OUTDIR/"cover_letter_CBC.docx"; doc.save(out); return out


def build_si():
    doc=Document(); default_style(doc); layout(doc)
    doc.add_heading("Supplementary Information", level=1); parse_md(doc,SI_MD,include_figures=True)
    out=OUTDIR/"supplementary_information_CBC.docx"; doc.save(out); return out


if __name__ == "__main__":
    for f in (build_main(), build_cover(), build_si()): print(f)
